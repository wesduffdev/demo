"""Tests for the reusable Excel workbook engine (BD-006, ``buyers_desk.reporting``).

Covers: sheet-per-``ReportSheet`` rendering, header text + styling, data-cell
fidelity, empty-``sheets`` handling, and that ``write_workbook`` produces a
real, reopenable ``.xlsx`` at the caller-supplied path (never a hardcoded
one). All sample data is synthetic SKU/brand/store data — no customer PII,
per the no-pii rule.
"""

from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path

import openpyxl
import pytest

from buyers_desk.contracts import Report, ReportKind, ReportSheet
from buyers_desk.reporting import (
    FROZEN_HEADER_CELL,
    HEADER_FILL,
    HEADER_FONT,
    build_workbook,
    write_workbook,
)
from buyers_desk.reporting.workbook import EMPTY_WORKBOOK_SHEET_NAME

FIXED_NOW = datetime(2026, 7, 19, 12, 0, 0, tzinfo=timezone.utc)


def _sample_report() -> Report:
    reorder_sheet = ReportSheet(
        name="Reorder",
        headers=["SKU", "Product", "On Hand", "Suggested Qty"],
        rows=[
            ["SKU-001", "Test Snack Bar", 12, 48],
            ["SKU-002", "Test Trail Mix", 3, 120],
        ],
    )
    slow_mover_sheet = ReportSheet(
        name="Slow Movers / Overstock",
        headers=["SKU", "Product", "Weeks On Hand"],
        rows=[["SKU-010", "Test Novelty Mug", 26.5]],
    )
    return Report(
        report_id="rpt-bd006-001",
        kind=ReportKind.LOW_STOCK_REORDER,
        generated_at=FIXED_NOW,
        sheets=[reorder_sheet, slow_mover_sheet],
    )


# ---------------------------------------------------------------------------
# build_workbook — pure builder, no disk I/O
# ---------------------------------------------------------------------------


def test_build_workbook_creates_one_sheet_per_report_sheet_in_order():
    report = _sample_report()
    workbook = build_workbook(report)

    assert workbook.sheetnames[0] == "Reorder"
    # "/" is invalid in an openpyxl sheet title, so it must be sanitized.
    assert "/" not in workbook.sheetnames[1]
    assert len(workbook.sheetnames) == 2


def test_build_workbook_does_not_touch_publish_fields():
    report = _sample_report()
    build_workbook(report)

    assert report.published_to is None
    assert report.published_at is None


def test_build_workbook_handles_empty_sheets_list_gracefully():
    report = Report(
        report_id="rpt-empty", kind=ReportKind.SALES_PERFORMANCE, generated_at=FIXED_NOW
    )

    workbook = build_workbook(report)

    assert len(workbook.sheetnames) == 1
    assert workbook.sheetnames[0] == EMPTY_WORKBOOK_SHEET_NAME


def test_build_workbook_deduplicates_sheet_titles_that_collide_after_sanitizing():
    report = Report(
        report_id="rpt-dupe",
        kind=ReportKind.SLOW_MOVER_OVERSTOCK,
        generated_at=FIXED_NOW,
        sheets=[
            ReportSheet(name="A/B", headers=["SKU"], rows=[["SKU-020"]]),
            ReportSheet(name="A:B", headers=["SKU"], rows=[["SKU-021"]]),
        ],
    )

    workbook = build_workbook(report)

    assert len(workbook.sheetnames) == 2
    assert len(set(workbook.sheetnames)) == 2


# ---------------------------------------------------------------------------
# write_workbook — writes a real, reopenable .xlsx to the caller-supplied path
# ---------------------------------------------------------------------------


def test_write_workbook_writes_to_the_caller_supplied_path(tmp_path):
    report = _sample_report()
    target = tmp_path / "nested" / "reorder-report.xlsx"

    result_path = write_workbook(report, target)

    assert result_path == target
    assert target.exists()
    assert isinstance(result_path, Path)


def test_write_workbook_reopened_sheet_names_and_headers_match(tmp_path):
    report = _sample_report()
    target = tmp_path / "report.xlsx"
    write_workbook(report, target)

    reopened = openpyxl.load_workbook(target)

    assert reopened.sheetnames[0] == "Reorder"
    reorder_ws = reopened["Reorder"]
    header_cells = [reorder_ws.cell(row=1, column=col).value for col in range(1, 5)]
    assert header_cells == ["SKU", "Product", "On Hand", "Suggested Qty"]


def test_write_workbook_reopened_known_data_cell_matches(tmp_path):
    report = _sample_report()
    target = tmp_path / "report.xlsx"
    write_workbook(report, target)

    reopened = openpyxl.load_workbook(target)
    reorder_ws = reopened["Reorder"]

    # Row 2 is the first data row (row 1 is the header).
    assert reorder_ws.cell(row=2, column=1).value == "SKU-001"
    assert reorder_ws.cell(row=2, column=3).value == 12
    assert reorder_ws.cell(row=3, column=1).value == "SKU-002"
    assert reorder_ws.cell(row=3, column=4).value == 120


def test_write_workbook_reopened_header_styling_was_actually_applied(tmp_path):
    report = _sample_report()
    target = tmp_path / "report.xlsx"
    write_workbook(report, target)

    reopened = openpyxl.load_workbook(target)
    reorder_ws = reopened["Reorder"]
    header_cell = reorder_ws.cell(row=1, column=1)

    assert header_cell.font.bold is True
    assert header_cell.font.color.rgb == HEADER_FONT.color.rgb
    assert header_cell.fill.fgColor.rgb == HEADER_FILL.fgColor.rgb
    assert reorder_ws.freeze_panes == FROZEN_HEADER_CELL


def test_write_workbook_applies_thousands_number_format_to_numeric_cells(tmp_path):
    report = _sample_report()
    target = tmp_path / "report.xlsx"
    write_workbook(report, target)

    reopened = openpyxl.load_workbook(target)
    reorder_ws = reopened["Reorder"]

    # "On Hand" (int) and "Suggested Qty" (int) columns get a thousands format.
    assert "#,##0" in reorder_ws.cell(row=2, column=3).number_format
    assert "#,##0" in reorder_ws.cell(row=2, column=4).number_format

    slow_mover_ws = reopened[reopened.sheetnames[1]]
    # "Weeks On Hand" (float) gets the decimal thousands format.
    assert "0.00" in slow_mover_ws.cell(row=2, column=3).number_format


def test_write_workbook_does_not_populate_publish_fields_on_the_report(tmp_path):
    report = _sample_report()
    write_workbook(report, tmp_path / "report.xlsx")

    assert report.published_to is None
    assert report.published_at is None


def test_write_workbook_column_widths_are_set_for_readability(tmp_path):
    report = _sample_report()
    target = tmp_path / "report.xlsx"
    write_workbook(report, target)

    reopened = openpyxl.load_workbook(target)
    reorder_ws = reopened["Reorder"]

    # Column A ("SKU") should be sized wide enough to show its content.
    assert reorder_ws.column_dimensions["A"].width >= len("SKU-001")


def test_workbook_module_does_not_import_a_sharepoint_publisher():
    import ast

    import buyers_desk.reporting.workbook as workbook_module

    module_source = Path(workbook_module.__file__).read_text(encoding="utf-8")
    tree = ast.parse(module_source)
    imported_names = set()
    for node in ast.walk(tree):
        if isinstance(node, ast.Import):
            imported_names.update(alias.name for alias in node.names)
        elif isinstance(node, ast.ImportFrom) and node.module:
            imported_names.add(node.module)

    assert not any("sharepoint" in name.lower() for name in imported_names)


@pytest.mark.parametrize("bad_char", ["/", "\\", "[", "]", "*", "?", ":"])
def test_build_workbook_sanitizes_every_invalid_sheet_title_character(bad_char):
    report = Report(
        report_id="rpt-char",
        kind=ReportKind.SALES_PERFORMANCE,
        generated_at=FIXED_NOW,
        sheets=[ReportSheet(name=f"Bad{bad_char}Name", headers=["SKU"], rows=[["SKU-030"]])],
    )

    workbook = build_workbook(report)

    assert bad_char not in workbook.sheetnames[0]
