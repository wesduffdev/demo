"""Generate a report workbook from local mock data — no credentials, no network.

This is the **mock end-to-end entry point**. It reads the local sample WAR
workbook (``data/samples/WAR_Executive_Overview_6.29.26.xlsx`` by default) and
renders it through the shared report engine
(:mod:`buyers_desk.reporting.workbook`) into an Excel-compatible file under
``./output/``. Run it with::

    python -m buyers_desk.mock_report            # -> output/WAR_<today>.xlsx
    make mock-report

Scope (honest about what it is)
-------------------------------
It is a stand-in for the full ``pull -> snapshot -> analyze -> report`` cycle
(BD-022, roadmap): it does **not** compute KPIs, flags, or recommendations. It
faithfully renders the mock source's sheets so you can exercise the
``mock data -> Report -> .xlsx`` path with no Fulfil connection and no
credentials. Row 1 of each source sheet is treated as the header row; the rest
are data rows, read with ``data_only=True`` (computed values, not formulas).

Data handling
-------------
The default mock workbook is **local-only / git-ignored** and holds real
(confidential, PII-clean) financials — see ``data/samples/README.md``. A
missing source file exits with a clear message (not a traceback), since a fresh
clone will not contain it. The generated file lands in ``./output/``, which is
also git-ignored (see ``output/README.md``).
"""

from __future__ import annotations

import argparse
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Optional, Sequence

from openpyxl import load_workbook

from buyers_desk.contracts import Report, ReportKind, ReportSheet
from buyers_desk.reporting.workbook import write_workbook

#: Default mock data source (local-only / git-ignored — see module docstring).
DEFAULT_SOURCE = "data/samples/WAR_Executive_Overview_6.29.26.xlsx"

#: Default delivery folder for the generated workbook (git-ignored).
DEFAULT_OUTPUT_DIR = "output"


def _today() -> str:
    """UTC date as ``YYYY-MM-DD`` — used for the default report date/filename."""
    return datetime.now(timezone.utc).strftime("%Y-%m-%d")


def load_mock_report(
    source: str | Path,
    *,
    report_date: Optional[str] = None,
    generated_at: Optional[datetime] = None,
) -> Report:
    """Read every sheet of the mock workbook at ``source`` into a :class:`Report`.

    Row 1 of each sheet becomes the header row (``None`` cells rendered as
    empty strings); the remaining rows are data rows. Values are read with
    ``data_only=True`` so cached computed values are used rather than formula
    text. Faithful render only — no figures are altered.
    """
    workbook = load_workbook(Path(source), read_only=True, data_only=True)
    try:
        sheets: list[ReportSheet] = []
        for worksheet in workbook.worksheets:
            rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
            raw_headers = rows[0] if rows else []
            headers = ["" if cell is None else str(cell) for cell in raw_headers]
            data_rows = rows[1:]
            sheets.append(ReportSheet(name=worksheet.title, headers=headers, rows=data_rows))
    finally:
        workbook.close()

    return Report(
        report_id=f"mock-{report_date or _today()}",
        kind=ReportKind.SALES_PERFORMANCE,
        generated_at=generated_at or datetime.now(timezone.utc),
        sheets=sheets,
    )


def generate(
    source: str | Path = DEFAULT_SOURCE,
    output_dir: str | Path = DEFAULT_OUTPUT_DIR,
    *,
    report_date: Optional[str] = None,
) -> Path:
    """Render the mock workbook at ``source`` to ``output_dir/WAR_<date>.xlsx``.

    Returns the path written. Creates ``output_dir`` if needed (via
    :func:`~buyers_desk.reporting.workbook.write_workbook`).
    """
    resolved_date = report_date or _today()
    report = load_mock_report(source, report_date=resolved_date)
    out_path = Path(output_dir) / f"WAR_{resolved_date}.xlsx"
    return write_workbook(report, out_path)


def main(argv: Optional[Sequence[str]] = None) -> int:
    """CLI entry point. Returns a process exit code (0 ok, 1 missing source)."""
    parser = argparse.ArgumentParser(
        prog="python -m buyers_desk.mock_report",
        description="Generate a report workbook from local mock data (no Fulfil, no credentials).",
    )
    parser.add_argument(
        "--source",
        default=DEFAULT_SOURCE,
        help=f"Mock data workbook to render (default: {DEFAULT_SOURCE}).",
    )
    parser.add_argument(
        "--output-dir",
        default=DEFAULT_OUTPUT_DIR,
        help=f"Folder to write the report into (default: {DEFAULT_OUTPUT_DIR}).",
    )
    parser.add_argument(
        "--report-date",
        default=None,
        help="Report date for the filename/id as YYYY-MM-DD (default: today, UTC).",
    )
    args = parser.parse_args(argv)

    source = Path(args.source)
    if not source.is_file():
        print(
            f"error: mock data file not found: {source}\n"
            "It is local-only / git-ignored — place the sample WAR workbook there, or pass "
            "--source <path>. See data/samples/README.md.",
            file=sys.stderr,
        )
        return 1

    out_path = generate(
        source=args.source,
        output_dir=args.output_dir,
        report_date=args.report_date,
    )
    print(f"wrote {out_path}")
    return 0


if __name__ == "__main__":  # pragma: no cover - exercised via the CLI, not tests
    raise SystemExit(main())
