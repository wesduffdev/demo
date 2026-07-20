"""Reusable Excel workbook builder for the ``Report`` aggregate (BD-006).

This module is the shared rendering engine that the three concrete v1
reports — sales performance (BD-017), low-stock/reorder (BD-018), and
slow-mover/overstock (BD-019) — will all render through, so every workbook
``report-publisher`` produces has the same header styling, column sizing,
frozen header row, and number formatting. It builds ONLY the workbook: it
never writes/knows about a SharePoint destination (that is BD-007) and never
alters the numbers or recommendations it is handed — it renders a ``Report``
faithfully (see the report-publisher agent's "must NOT" list).

Public API
----------
``build_workbook(report)`` builds an in-memory ``openpyxl.Workbook`` (pure —
no disk I/O). ``write_workbook(report, path)`` builds and saves it to a
caller-supplied path, returning that path as a ``Path``. Neither function
populates ``Report.published_to``/``published_at`` — those stay ``None``
until a human approves the publish step (human-in-the-loop rule) and
BD-007's publisher sets them explicitly.

Reusable styling primitives
----------------------------
``HEADER_FONT``, ``HEADER_FILL``, and ``HEADER_ALIGNMENT`` are the shared
header look; ``_style_header_row`` applies them plus a header row height.
``_autosize_columns`` sizes every column from its header/content width
(clamped to ``MIN_COLUMN_WIDTH``/``MAX_COLUMN_WIDTH``). ``FROZEN_HEADER_CELL``
is the ``freeze_panes`` target that keeps the header row visible while
scrolling. ``THOUSANDS_NUMBER_FORMAT``/``THOUSANDS_DECIMAL_NUMBER_FORMAT`` are
applied automatically to ``int``/``float`` data cells respectively (see
``_apply_number_formats``) so every quantity/currency column reads with
thousands separators without each report needing to know openpyxl's number
format syntax. All of this lives in one module precisely so BD-017/018/019
share it instead of re-implementing formatting per report.

Empty-``sheets`` behavior
--------------------------
openpyxl requires every workbook to have at least one worksheet.
``build_workbook`` on a ``Report`` with an empty ``sheets`` list does NOT
raise — it returns a single placeholder worksheet named
``EMPTY_WORKBOOK_SHEET_NAME`` with no headers/rows, so an otherwise-valid
(if data-less) ``Report`` still produces a valid, openable ``.xlsx`` rather
than failing the whole publish attempt. Callers that want an empty report to
be an error should check ``report.sheets`` before calling.
"""

from __future__ import annotations

import os
from collections.abc import Sequence
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from buyers_desk.contracts import Report, ReportSheet

# ---------------------------------------------------------------------------
# Reusable styling primitives — shared by every report built through this
# engine (BD-017/018/019). Keep changes here, not per-report, so formatting
# stays consistent across the whole reporting surface.
# ---------------------------------------------------------------------------

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="FF1F4E78", end_color="FF1F4E78", fill_type="solid")
HEADER_ALIGNMENT = Alignment(horizontal="left", vertical="center", wrap_text=False)
HEADER_ROW_HEIGHT = 20

# ``freeze_panes`` target that pins row 1 (the header) while the sheet scrolls.
FROZEN_HEADER_CELL = "A2"

MIN_COLUMN_WIDTH = 8
MAX_COLUMN_WIDTH = 60
COLUMN_WIDTH_PADDING = 2

# Applied to int/float data cells respectively so quantities/currency read
# with thousands separators everywhere, consistently.
THOUSANDS_NUMBER_FORMAT = "#,##0"
THOUSANDS_DECIMAL_NUMBER_FORMAT = "#,##0.00"

# Used when a ``Report`` has no sheets at all (see module docstring).
EMPTY_WORKBOOK_SHEET_NAME = "Empty Report"

# openpyxl forbids these characters in a sheet title and caps the length at 31.
_INVALID_SHEET_TITLE_CHARS = frozenset("[]:*?/\\")
_MAX_SHEET_TITLE_LENGTH = 31


def _safe_sheet_title(name: str) -> str:
    """Sanitize a ``ReportSheet.name`` into a valid openpyxl sheet title."""
    cleaned = "".join("-" if ch in _INVALID_SHEET_TITLE_CHARS else ch for ch in name.strip())
    cleaned = cleaned or "Sheet"
    return cleaned[:_MAX_SHEET_TITLE_LENGTH]


def _unique_sheet_title(title: str, used_titles: set[str]) -> str:
    """Disambiguate a sheet title against ones already used in this workbook."""
    if title not in used_titles:
        return title
    suffix = 2
    while True:
        marker = f" ({suffix})"
        candidate = f"{title[: _MAX_SHEET_TITLE_LENGTH - len(marker)]}{marker}"
        if candidate not in used_titles:
            return candidate
        suffix += 1


def _style_header_row(worksheet: Worksheet, column_count: int) -> None:
    """Apply the shared header look (bold white on dark fill) to row 1."""
    for col_idx in range(1, column_count + 1):
        cell = worksheet.cell(row=1, column=col_idx)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGNMENT
    worksheet.row_dimensions[1].height = HEADER_ROW_HEIGHT


def _apply_number_formats(worksheet: Worksheet, column_count: int, row_count: int) -> None:
    """Give every numeric data cell a consistent thousands-separator format.

    Data occupies rows 2..row_count+1 (row 1 is the header). ``bool`` is
    checked first since ``bool`` is a subclass of ``int`` in Python.
    """
    for row_idx in range(2, row_count + 2):
        for col_idx in range(1, column_count + 1):
            cell = worksheet.cell(row=row_idx, column=col_idx)
            value = cell.value
            if isinstance(value, bool) or value is None:
                continue
            if isinstance(value, int):
                cell.number_format = THOUSANDS_NUMBER_FORMAT
            elif isinstance(value, float):
                cell.number_format = THOUSANDS_DECIMAL_NUMBER_FORMAT


def _autosize_columns(
    worksheet: Worksheet, headers: Sequence[str], rows: Sequence[Sequence[object]]
) -> None:
    """Size every column from its header/content width, clamped to sane bounds."""
    column_count = max(len(headers), max((len(row) for row in rows), default=0))
    for col_idx in range(1, column_count + 1):
        header_text = headers[col_idx - 1] if col_idx <= len(headers) else ""
        widest = len(str(header_text))
        for row in rows:
            if col_idx <= len(row) and row[col_idx - 1] is not None:
                widest = max(widest, len(str(row[col_idx - 1])))
        width = min(MAX_COLUMN_WIDTH, max(MIN_COLUMN_WIDTH, widest + COLUMN_WIDTH_PADDING))
        worksheet.column_dimensions[get_column_letter(col_idx)].width = width


def _render_sheet(worksheet: Worksheet, sheet: ReportSheet) -> None:
    """Write one ``ReportSheet`` (header + data rows) faithfully, then style it."""
    headers = list(sheet.headers)
    worksheet.append(headers)
    for row in sheet.rows:
        worksheet.append(list(row))

    if headers:
        _style_header_row(worksheet, len(headers))
        worksheet.freeze_panes = FROZEN_HEADER_CELL

    _apply_number_formats(worksheet, column_count=len(headers), row_count=len(sheet.rows))
    _autosize_columns(worksheet, headers=headers, rows=sheet.rows)


def build_workbook(report: Report) -> Workbook:
    """Render a ``Report`` into an in-memory ``openpyxl.Workbook`` (pure builder).

    One worksheet per ``ReportSheet``, in order, with the shared header
    styling/freeze/number-format/column-sizing primitives applied. Renders
    the numbers and recommendations exactly as given — never alters them.
    Does not write to disk and does not touch ``published_to``/``published_at``.

    If ``report.sheets`` is empty, returns a workbook with a single
    placeholder sheet (see the module docstring) rather than raising, since
    openpyxl requires at least one worksheet.
    """
    workbook = Workbook()
    default_sheet = workbook.active

    if not report.sheets:
        default_sheet.title = _safe_sheet_title(EMPTY_WORKBOOK_SHEET_NAME)
        return workbook

    used_titles: set[str] = set()
    for index, sheet in enumerate(report.sheets):
        title = _unique_sheet_title(_safe_sheet_title(sheet.name), used_titles)
        used_titles.add(title)

        worksheet = default_sheet if index == 0 else workbook.create_sheet()
        worksheet.title = title
        _render_sheet(worksheet, sheet)

    return workbook


def write_workbook(report: Report, path: str | os.PathLike) -> Path:
    """Build ``report`` and save it as a valid ``.xlsx`` at ``path``.

    ``path`` is caller-supplied — this module never hardcodes a destination
    directory (SharePoint delivery is a later ticket; see BD-007). Creates
    any missing parent directories, then returns ``path`` as a ``Path``.
    """
    workbook = build_workbook(report)
    output_path = Path(path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return output_path
